"""Field Mode: photo (+ voice note) per observation -> Claude vision -> structured entry; visit -> report.

A visit is one site walk. Each observation is processed in the background as it's uploaded, so the
user keeps capturing. "Compile report" builds a DOCX with photos, a Notion page, tasks for problems,
and a Slack post.
"""
import base64
import io
import traceback
from datetime import date
from pathlib import Path

from fastapi import APIRouter, BackgroundTasks, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from PIL import Image, ImageOps
from pydantic import BaseModel
from sqlmodel import select

import notion
import slack
from api import current_user
from brain import CLAUDE_MODEL, EXTRA, client
from config import DATA, DRY_RUN
from db import FieldVisit, Observation, Task, User, session
from pipeline import match_user
from transcribe import transcribe

router = APIRouter(prefix="/field")
FIELD_DIR = DATA / "field"
SEVERITY_RANK = {"high": 3, "medium": 2, "low": 1, "none": 0}

OBSERVE_TOOL = {
    "name": "record_observation",
    "description": "Record one field observation from the photo and the inspector's voice note.",
    "input_schema": {
        "type": "object",
        "additionalProperties": False,
        "properties": {
            "title": {"type": "string", "description": "3-8 word English title, e.g. 'Cracked floor tile near lift'."},
            "observation": {"type": "string", "description": "1-3 English sentences: what the photo shows, using the voice note for context."},
            "category": {"type": "string", "enum": ["structural", "electrical", "plumbing", "safety", "finishing", "cleanliness", "equipment", "inventory", "other"]},
            "severity": {"type": "string", "enum": ["none", "low", "medium", "high"], "description": "none if nothing is wrong."},
            "action": {"type": "string", "description": "Recommended next step in English, or empty if none."},
            "owner": {"type": "string", "description": "Person the voice note assigns the action to, or empty."},
        },
        "required": ["title", "observation", "category", "severity", "action", "owner"],
    },
}

REPORT_TOOL = {
    "name": "record_report",
    "description": "Record the site visit report.",
    "input_schema": {
        "type": "object",
        "additionalProperties": False,
        "properties": {
            "summary": {"type": "string", "description": "2-3 short English sentences a manager can read on a phone."},
            "key_issues": {"type": "array", "items": {"type": "string"}, "description": "Up to 5 problems, worst first, each under 15 words."},
            "recommendations": {"type": "array", "items": {"type": "string"}, "description": "Up to 5 concrete next steps, each under 15 words."},
        },
        "required": ["summary", "key_issues", "recommendations"],
    },
}


def _jpeg_b64(path: Path, max_side: int = 1568) -> str:
    """Phone photos are big; Claude reads 1568 px fine. Also fixes EXIF rotation."""
    img = ImageOps.exif_transpose(Image.open(path)).convert("RGB")
    img.thumbnail((max_side, max_side))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=85)
    return base64.b64encode(buf.getvalue()).decode()


def _tool_input(content: list[dict], tool_name: str, max_tokens: int = 4000) -> dict:
    msg = client.beta.messages.create(
        model=CLAUDE_MODEL,
        max_tokens=max_tokens,
        tools=[OBSERVE_TOOL if tool_name == "record_observation" else REPORT_TOOL],
        messages=[{"role": "user", "content": content}],
        **EXTRA,
    )
    block = next((b for b in msg.content if b.type == "tool_use" and b.name == tool_name), None)
    if block is None:
        raise RuntimeError(f"Claude didn't call {tool_name} (stop_reason={msg.stop_reason})")
    return block.input


def observe(photo: Path, voice_text: str, note: str, site: str) -> dict:
    context = f"Site: {site}. Today: {date.today():%A %Y-%m-%d}."
    if voice_text:
        context += f"\nInspector's voice note (may be Telugu/Hindi/English): {voice_text}"
    if note:
        context += f"\nInspector's typed note: {note}"
    return _tool_input(
        [
            {"type": "image", "source": {"type": "base64", "media_type": "image/jpeg", "data": _jpeg_b64(photo)}},
            {"type": "text", "text": context + "\n\nLook at the photo and call record_observation once. "
             "Describe only what is visible or stated; don't guess causes you can't see."},
        ],
        "record_observation",
    )


def process_observation(obs_id: int) -> None:
    with session() as s:
        o = s.get(Observation, obs_id)
        visit = s.get(FieldVisit, o.visit_id)
        try:
            if o.voice_path:
                o.voice_text = transcribe(Path(o.voice_path))
            result = observe(Path(o.photo_path), o.voice_text, o.note, visit.site)
            for k in ("title", "observation", "category", "severity", "action", "owner"):
                setattr(o, k, result.get(k, ""))
            o.status = "done"
        except Exception as e:
            traceback.print_exc()
            o.status, o.error = "error", str(e)[:500]
        s.add(o)
        s.commit()


def _docx(visit: FieldVisit, obs: list[Observation]) -> Path:
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    doc.add_heading(f"Site report: {visit.site}", 0)
    doc.add_paragraph(f"Date: {visit.created_at:%d %b %Y}  ·  Observations: {len(obs)}")
    doc.add_heading("Summary", 1)
    doc.add_paragraph(visit.summary)
    for heading, text in (("Key issues", visit.key_issues), ("Recommendations", visit.recommendations)):
        if text:
            doc.add_heading(heading, 1)
            for line in text.splitlines():
                doc.add_paragraph(line, style="List Bullet")
    doc.add_heading("Observations", 1)
    for i, o in enumerate(obs, 1):
        doc.add_heading(f"{i}. {o.title}  [{o.severity or 'n/a'}]", 2)
        buf = io.BytesIO()
        img = ImageOps.exif_transpose(Image.open(o.photo_path)).convert("RGB")
        img.thumbnail((1200, 1200))
        img.save(buf, "JPEG", quality=80)
        buf.seek(0)
        doc.add_picture(buf, width=Inches(4.5))
        doc.add_paragraph(o.observation)
        if o.action:
            doc.add_paragraph(f"Action: {o.action}" + (f"  (owner: {o.owner})" if o.owner else ""))
        if o.voice_text:
            doc.add_paragraph().add_run(f"Voice note: {o.voice_text}").italic = True
    out = FIELD_DIR / f"visit-{visit.id}" / f"report-{visit.id}.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


def build_report(visit_id: int) -> None:
    with session() as s:
        visit = s.get(FieldVisit, visit_id)
        obs = list(s.exec(select(Observation).where(Observation.visit_id == visit_id, Observation.status == "done").order_by(Observation.id)))
        users = list(s.exec(select(User).where(User.team_id == visit.team_id)))
        try:
            listing = "\n".join(
                f"{i}. [{o.severity}] {o.title}: {o.observation} Action: {o.action or '-'}" for i, o in enumerate(obs, 1)
            )
            rep = _tool_input(
                [{"type": "text", "text": f"Site visit to {visit.site}, {len(obs)} observations:\n{listing}\n\n"
                  "Write a short, phone-friendly report and call record_report once."}],
                "record_report",
            )
            visit.summary = rep["summary"]
            visit.key_issues = "\n".join(rep["key_issues"])
            visit.recommendations = "\n".join(rep["recommendations"])

            # A task for every problem that needs doing (medium/high, or anything with an owner).
            for o in obs:
                if o.task_id or not o.action or (SEVERITY_RANK.get(o.severity, 0) < 2 and not o.owner):
                    continue
                user = match_user(o.owner, users) if o.owner else None
                owner = user.name if user else (o.owner or "unassigned")
                priority = "high" if o.severity == "high" else "medium"
                url = "" if DRY_RUN else notion.create_task(title=o.action, owner=owner, due=None, priority=priority, source="cloud")
                t = Task(team_id=visit.team_id, owner_name=owner, owner_user_id=user and user.id, title=o.action, priority=priority, notion_url=url)
                s.add(t)
                s.flush()
                o.task_id = t.id
                s.add(o)

            if not DRY_RUN:
                visit.notion_url = notion.create_note(
                    title=f"Site report: {visit.site}",
                    summary=visit.summary,
                    decisions=[*rep["key_issues"], *[f"Recommendation: {r}" for r in rep["recommendations"]]],
                    transcript="\n".join(f"[{o.severity}] {o.title}: {o.observation}" for o in obs),
                )
            visit.report_path = str(_docx(visit, obs))
            if slack.enabled():
                try:
                    slack.post(
                        slack.meeting_blocks(f"Site report: {visit.site}", visit.summary, [], [], rep["key_issues"][:5], visit.notion_url),
                        f"Site report: {visit.site}",
                    )
                except Exception:
                    pass
            visit.report_status = "done"
            s.add(visit)
            s.commit()
        except Exception as e:
            traceback.print_exc()
            s.rollback()
            visit = s.get(FieldVisit, visit_id)
            visit.report_status, visit.report_error = "error", str(e)[:500]
            s.add(visit)
            s.commit()


# ---------- API ----------

class NewVisit(BaseModel):
    site: str


def _visit(s, visit_id: int, user: User) -> FieldVisit:
    v = s.get(FieldVisit, visit_id)
    if not v or v.team_id != user.team_id:
        raise HTTPException(404, "Visit not found")
    return v


def _obs_out(o: Observation) -> dict:
    out = o.model_dump(exclude={"photo_path", "voice_path", "visit_id"})
    out["photo_url"] = f"/field/photos/{o.id}"
    return out


@router.post("/visits")
def new_visit(body: NewVisit, user: User = Depends(current_user)):
    with session() as s:
        v = FieldVisit(team_id=user.team_id, created_by=user.id, site=body.site.strip())
        s.add(v)
        s.commit()
        s.refresh(v)
        return {"id": v.id, "site": v.site}


@router.get("/visits")
def list_visits(user: User = Depends(current_user)):
    with session() as s:
        rows = s.exec(select(FieldVisit).where(FieldVisit.team_id == user.team_id).order_by(FieldVisit.created_at.desc())).all()
        return [{"id": v.id, "site": v.site, "report_status": v.report_status, "created_at": v.created_at} for v in rows]


@router.get("/visits/{visit_id}")
def get_visit(visit_id: int, user: User = Depends(current_user)):
    with session() as s:
        v = _visit(s, visit_id, user)
        obs = s.exec(select(Observation).where(Observation.visit_id == v.id).order_by(Observation.id)).all()
        return {
            **v.model_dump(exclude={"report_path", "team_id"}),
            "key_issues": [x for x in v.key_issues.splitlines() if x],
            "recommendations": [x for x in v.recommendations.splitlines() if x],
            "report_url": f"/field/visits/{v.id}/report.docx" if v.report_path else "",
            "observations": [_obs_out(o) for o in obs],
        }


@router.post("/visits/{visit_id}/observations")
async def add_observation(
    visit_id: int,
    background: BackgroundTasks,
    photo: UploadFile = File(...),
    voice: UploadFile | None = File(None),
    note: str = Form(""),
    user: User = Depends(current_user),
):
    """One photo (+ optional voice note / typed note). Returns at once; poll GET /field/visits/{id}."""
    with session() as s:
        v = _visit(s, visit_id, user)
        folder = FIELD_DIR / f"visit-{v.id}"
        folder.mkdir(parents=True, exist_ok=True)
        o = Observation(visit_id=v.id, photo_path="", note=note)
        s.add(o)
        s.commit()
        s.refresh(o)
        photo_path = folder / f"obs-{o.id}{Path(photo.filename or '').suffix or '.jpg'}"
        photo_path.write_bytes(await photo.read())
        o.photo_path = str(photo_path)
        if voice is not None and voice.filename:
            voice_path = folder / f"obs-{o.id}-voice{Path(voice.filename).suffix or '.m4a'}"
            voice_path.write_bytes(await voice.read())
            o.voice_path = str(voice_path)
        s.add(o)
        s.commit()
        obs_id = o.id
    background.add_task(process_observation, obs_id)
    return {"id": obs_id, "status": "processing"}


@router.get("/photos/{obs_id}")
def photo(obs_id: int, user: User = Depends(current_user)):
    with session() as s:
        o = s.get(Observation, obs_id)
        if not o:
            raise HTTPException(404, "Photo not found")
        _visit(s, o.visit_id, user)
        return FileResponse(o.photo_path)


@router.post("/visits/{visit_id}/report")
def compile_report(visit_id: int, background: BackgroundTasks, user: User = Depends(current_user)):
    """Build the report from all finished observations. Poll GET /field/visits/{id} for report_status."""
    with session() as s:
        v = _visit(s, visit_id, user)
        if v.report_status == "building":
            return {"report_status": "building"}
        v.report_status, v.report_error = "building", ""
        s.add(v)
        s.commit()
    background.add_task(build_report, visit_id)
    return {"report_status": "building"}


@router.get("/visits/{visit_id}/report.docx")
def report_docx(visit_id: int, user: User = Depends(current_user)):
    with session() as s:
        v = _visit(s, visit_id, user)
        if not v.report_path:
            raise HTTPException(404, "Report not built yet")
        return FileResponse(v.report_path, filename=f"site-report-{v.id}.docx")
